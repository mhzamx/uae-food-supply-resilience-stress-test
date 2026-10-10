"""P2 UAE Food Supply Resilience Stress Test - analytical model tables (gap method v0.1R).

Builds the English tables used in results/ and in the Power BI model from:
  - data/model_inputs_2026-10-06/  values frozen from the project workbook (see each file's source column)
  - outputs/                        corridor series produced by src/p2_corridors_v0.2.py

Segment requirements, corridor increments and the C1-only Low gaps are the values REGISTERED in the
project workbook (sheet "Brecha v0.1", sections I, J, K, L, Q, R). Subgroup requirements inside the
container segment are split with the registered section-K proportions, so every gap reproduces the
workbook exactly. A reconstruction from the published rounded inputs (RNI x k_ship x physical mix) is
reported in model_checks.csv as a consistency check.

Nothing here is observed food coverage: C_phys is a physical ceiling and C_ref a reference-allocation
benchmark (HS 01-24 share of 2023 non-fuel import weight, confidence C).

Usage (from the repository root):
  python src/p2_corridors_v0.2.py --data-dir data/snapshot_2026-10-05 --out-dir outputs
  python src/p2_build_model.py --inputs data/model_inputs_2026-10-06 --outputs outputs --results results [--xlsx powerbi/P2_PowerBI_model.xlsx]
"""
import argparse, os
import numpy as np
import pandas as pd

DEMANDS = ["Low", "High"]
SCOPES = ["C1-only", "Regional envelope", "C1 total (sensitivity)"]
ACTIVITIES = ["Low", "Base", "High synchronized"]
PORT_SEGMENTS = ["CONT", "GS"]
CORRIDOR_NAMES = {"C1": "C1 UAE east coast (Fujairah, Khor Fakkan)", "C2a": "C2a Oman: Sohar", "C2b": "C2b Oman: Duqm, Salalah",
                  "C3a": "C3a Saudi Red Sea: Jeddah, KAP (low reliability)"}
SEGMENT_NAMES = {"CONT": "Container", "GS": "Dry bulk", "CG_RORO": "General cargo + RoRo",
                 "TER": "Road (regional origin, not observed)", "GL": "Liquid bulk (not estimable)"}


def load(inputs):
    r = lambda f: pd.read_csv(os.path.join(inputs, f))
    reg = r("registered_values.csv").set_index("key")["value"].astype(float)
    return dict(cat=r("categories.csv"), sub=r("subgroups.csv"), reg=reg, k=r("registered_gap_c1_low.csv"),
                inc=r("corridor_increments.csv"), sup=r("suppliers_top5.csv"), route=r("route_exposure_top5.csv"),
                nut=r("nutrition_2023.csv"), const=r("constants.csv").set_index("key")["value"].astype(float))


def subgroup_requirements(d):
    """Shipment requirement by subgroup and demand level (t/day), consistent with the registered totals."""
    cat, sub, reg, k = d["cat"].set_index("category_id"), d["sub"].set_index("subgroup_id"), d["reg"], d["k"].set_index("subgroup_id")
    rows = []
    for dem in DEMANDS:
        col = f"gap_ref_{dem.lower()}_demand_t_day"
        cont_total = reg[f"req_CONT_{dem.lower()}"]
        # CONT: split with registered section-K proportions (exact reproduction of the workbook)
        for sid, g in k[col].items():
            rows.append((sid, dem, cont_total * g / k[col].sum()))
        # GS: wheat grain and raw sugar from published inputs, scaled to the registered GS total
        rni = lambda c: cat.loc[c, f"rni_a_2026_{dem.lower()}_t_day"]
        gs = {"wheat_grain": rni("wheat") * cat.loc["wheat", "k_ship"],
              "sugar_raw": rni("sugar") * cat.loc["sugar", "k_ship"] * sub.loc["sugar_raw", "mix_share_pct"] / 100}
        f = reg[f"req_GS_{dem.lower()}"] / sum(gs.values())
        rows += [(sid, dem, v * f) for sid, v in gs.items()]
        # TER: registered section-R subgroup values, scaled to the registered TER total
        ter = {"poultry_fresh": reg[f"ter_req_poultry_fresh_{dem.lower()}"], "dairy_fresh": reg[f"ter_req_dairy_fresh_{dem.lower()}"]}
        f = reg[f"req_TER_{dem.lower()}"] / sum(ter.values())
        rows += [(sid, dem, v * f) for sid, v in ter.items()]
        rows.append(("oil_crude", dem, reg[f"req_GL_{dem.lower()}"]))
    out = pd.DataFrame(rows, columns=["subgroup_id", "demand", "shipment_requirement_t_day"])
    return out.merge(sub[["category_id", "subgroup", "segment", "container_type"]], left_on="subgroup_id", right_index=True)


def reconstruction_check(d, req):
    """Rebuild segment totals from published rounded inputs and compare with the registered values."""
    cat, sub = d["cat"].set_index("category_id"), d["sub"]
    rows = []
    for dem in DEMANDS:
        rec = {}
        for _, s in sub[sub.segment.isin(["CONT", "GS", "TER", "GL"])].iterrows():
            c = cat.loc[s.category_id]
            v = c[f"rni_a_2026_{dem.lower()}_t_day"] * c.k_ship * s.mix_share_pct / 100
            rec[s.segment] = rec.get(s.segment, 0) + v
        for seg, v in rec.items():
            regv = d["reg"][f"req_{seg}_{dem.lower()}"]
            rows.append({"check": f"Segment requirement {seg}, {dem.lower()} demand", "registered": regv,
                         "rebuilt_from_rounded_inputs": round(v, 3), "difference": round(v - regv, 3),
                         "relative_difference_pct": round(100 * (v - regv) / regv, 4)})
    return pd.DataFrame(rows)


def scenarios(d):
    rows, i = [], 0
    for sc in SCOPES:
        for act in ACTIVITIES:
            i += 1
            rows.append({"scenario_id": f"S{i:02d}", "scope": sc, "activity_level": act,
                         "scenario": f"{sc} | {act}", "sort_order": i, "scope_order": SCOPES.index(sc) + 1,
                         "activity_order": ACTIVITIES.index(act) + 1, "is_sensitivity": sc.startswith("C1 total")})
    return pd.DataFrame(rows)


def coverage(d, scen, b):
    inc = d["inc"].merge(scen, on=["scope", "activity_level"])
    rows = []
    for _, r in inc[inc.segment.isin(PORT_SEGMENTS)].iterrows():
        for dem in DEMANDS:
            req = d["reg"][f"req_{r.segment}_{dem.lower()}"]
            delta = r.delta_cargo_t_day
            c_phys, c_ref = min(1.0, delta / req), min(1.0, delta * b / req)
            rows.append({"scenario_id": r.scenario_id, "scope": r.scope, "activity_level": r.activity_level, "demand": dem,
                         "segment": r.segment, "segment_name": SEGMENT_NAMES[r.segment],
                         "requirement_t_day": req, "delta_cargo_t_day": delta, "benchmark_food_share": b,
                         "c_phys": c_phys, "c_ref": c_ref, "gap_phys_t_day": req * (1 - c_phys), "gap_ref_t_day": req * (1 - c_ref),
                         "breakeven_food_share": req / delta, "cargo_needed_for_full_c_ref_t_day": req / b,
                         "food_equivalent_spare_or_shortfall_t_day": delta * b - req})
    return pd.DataFrame(rows)


def gaps_by_subgroup(req, cov, reg):
    payload = {"Dry": reg["payload_dry_40ft"], "Reefer": reg["payload_reefer_40ft_hc"]}
    m = req[req.segment.isin(PORT_SEGMENTS)].merge(
        cov[["scenario_id", "scope", "activity_level", "demand", "segment", "requirement_t_day", "delta_cargo_t_day", "c_phys", "c_ref"]]
        .rename(columns={"requirement_t_day": "segment_requirement_t_day", "delta_cargo_t_day": "segment_delta_cargo_t_day"}),
        on=["demand", "segment"])
    m["gap_ref_t_day"] = m.shipment_requirement_t_day * (1 - m.c_ref)
    m["gap_phys_t_day"] = m.shipment_requirement_t_day * (1 - m.c_phys)
    m["covered_ref_t_day"] = m.shipment_requirement_t_day * m.c_ref
    m["payload_t"] = m.container_type.map(payload)
    m["truck_equiv_trips_covered_lower_bound"] = m.covered_ref_t_day / m.payload_t
    m["truck_equiv_trips_gap_lower_bound"] = m.gap_ref_t_day / m.payload_t
    return m


def off_port(req, reg):
    share = {"poultry_fresh": reg["ter_share_poultry_fresh"], "dairy_fresh": reg["ter_share_dairy_fresh"], "oil_crude": reg["gl_regional_share"]}
    status = {"TER": "Regional-route requirement, not verified (road flow not observed)",
              "GL": "Port coverage not estimable (PortWatch liquid bulk dominated by petroleum)"}
    o = req[req.segment.isin(["TER", "GL"])].copy()
    o["regional_origin_share_upper_bound"] = o.subgroup_id.map(share)
    o["residual_after_upper_bound_t_day"] = o.shipment_requirement_t_day * (1 - o.regional_origin_share_upper_bound)
    o["status"] = o.segment.map(status)
    return o


def food_share_sensitivity(cov):
    grid = np.round(np.arange(0.05, 0.3501, 0.005), 3)
    rows = []
    for _, r in cov.iterrows():
        for s in grid:
            c = min(1.0, r.delta_cargo_t_day * s / r.requirement_t_day)
            rows.append({"scenario_id": r.scenario_id, "scope": r.scope, "activity_level": r.activity_level, "demand": r.demand,
                         "segment": r.segment, "food_share": s, "c_ref": c, "gap_ref_t_day": r.requirement_t_day * (1 - c)})
    return pd.DataFrame(rows)


def nutrition(d):
    cat, nut, pop = d["cat"].set_index("category_id"), d["nut"], d["const"]["fbs_population_2023"] * 1e6
    n = nut.copy()
    n["kcal_per_kg"] = n.kcal_per_capita_day_2023 * pop * 365 / (n.food_2023_kt * 1e6)
    n["protein_g_per_kg"] = n.protein_g_per_capita_day_2023 * pop * 365 / (n.food_2023_kt * 1e6)
    for dem in DEMANDS:
        rni = n.category_id.map(cat[f"rni_a_2026_{dem.lower()}_t_day"])
        n[f"import_dependent_million_kcal_day_{dem.lower()}"] = rni * 1000 * n.kcal_per_kg / 1e6
        n[f"import_dependent_protein_t_day_{dem.lower()}"] = rni * 1000 * n.protein_g_per_kg / 1e6
    for dem in DEMANDS:
        n[f"kcal_share_{dem.lower()}"] = n[f"import_dependent_million_kcal_day_{dem.lower()}"] / n[f"import_dependent_million_kcal_day_{dem.lower()}"].sum()
        n[f"protein_share_{dem.lower()}"] = n[f"import_dependent_protein_t_day_{dem.lower()}"] / n[f"import_dependent_protein_t_day_{dem.lower()}"].sum()
    return n


def suppliers(d):
    s = d["sup"]
    top = s[s["rank"] == 1][["supplier_group", "supplier_group_name", "n_suppliers", "hhi", "country", "share_value_pct", "top5_share_pct"]]
    top = top.rename(columns={"country": "top1_country", "share_value_pct": "top1_share_pct"})
    top3 = s[s["rank"] <= 3].groupby("supplier_group").share_value_pct.sum().rename("top3_share_pct")
    top = top.merge(top3, left_on="supplier_group", right_index=True).merge(d["route"], on="supplier_group", how="left")
    return top, s


def corridor_weekly(outputs):
    w = pd.read_csv(os.path.join(outputs, "corridor_ports_weekly.csv"), parse_dates=["week"])
    w["import_general_cargo_roro"] = w.import_general_cargo + w.import_roro
    cols = {"import_container": "CONT", "import_dry_bulk": "GS", "import_general_cargo_roro": "CG_RORO", "import_tanker": "TANKER"}
    g = w.groupby(["corridor", "week"], as_index=False)[list(cols)].sum()
    long = g.melt(id_vars=["corridor", "week"], value_vars=list(cols), var_name="col", value_name="imports_t_week")
    long["segment"] = long.col.map(cols); long = long.drop(columns="col")
    base = long[(long.week >= "2025-03-01") & (long.week <= "2026-02-27")].groupby(["corridor", "segment"]).imports_t_week.mean().rename("base_mean_t_week")
    long = long.merge(base, left_on=["corridor", "segment"], right_index=True)
    long["period"] = np.where(long.week >= "2026-03-01", "Closure", np.where(long.week >= "2025-03-01", "Pre-war base", "Earlier"))
    long["corridor_code"] = long.corridor.str.split(" ").str[0]
    long["corridor"] = long.corridor_code.map(CORRIDOR_NAMES)
    long["corridor_order"] = long.corridor_code.map({c: i for i, c in enumerate(CORRIDOR_NAMES, start=1)})
    return long.sort_values(["corridor", "segment", "week"])


README_ROWS = [
    ("UAE Food Supply Resilience Stress Test - Power BI data model (v0.1)",),
    ("Built by src/p2_build_model.py from data/model_inputs_2026-10-06 and outputs/ (p2_corridors_v0.2.py). Values only; no formulas. Units: t/day unless stated.",),
    ("C_phys = physical ceiling; C_ref = reference-allocation benchmark (food share 19.6656%, confidence C). Neither is observed food coverage.",),
    ("",),
    ("Table", "Grain (one row per)", "Keys / relationships", "Use"),
    ("commodities", "Physical subgroup", "subgroup_id (PK); category_id -> categories; supplier_group -> supplier_concentration", "Subgroup, segment and container type"),
    ("categories", "Basket category", "category_id (PK)", "RNI A/B 2026, k_ship, regional-origin share, priority, confidence"),
    ("demand_levels", "Demand level", "demand (PK)", "Slicer: Low / High 2026 resident demand"),
    ("segments", "Segment", "segment (PK)", "Segment names; port-model flag"),
    ("scenarios", "Scope x activity level", "scenario_id (PK)", "Slicer: C1-only, Regional envelope, C1 total (sensitivity) x Low/Base/High synchronized"),
    ("requirement_by_subgroup", "Subgroup x demand", "subgroup_id, demand", "Shipment requirement (registered totals)"),
    ("corridor_increments", "Scenario x segment", "scenario_id, segment", "Incremental total-cargo activity over the pre-war base"),
    ("coverage_by_segment", "Scenario x demand x segment (CONT, GS)", "scenario_id, demand, segment", "Requirement, delta cargo, C_phys, C_ref, gaps, break-even share, threshold"),
    ("gap_by_subgroup", "Scenario x demand x subgroup", "scenario_id, demand, subgroup_id", "Gap_ref/Gap_phys by subgroup; truck-equivalent lower bounds"),
    ("off_port_requirements", "Subgroup (TER, GL) x demand", "subgroup_id, demand", "Requirements outside the port model and regional-origin upper bound"),
    ("food_share_sensitivity", "Scenario x demand x segment x food share", "scenario_id, demand, segment", "C_ref and Gap_ref for food shares 5%-35%"),
    ("nutrition", "Category", "category_id", "kcal and protein linked to the import-dependent requirement"),
    ("supplier_concentration", "Supplier group", "supplier_group (PK)", "HHI, Top1/Top3/Top5, Red Sea exposure within Top5 (by value)"),
    ("suppliers_top5", "Supplier group x rank", "supplier_group", "Top 5 supplier countries and value shares"),
    ("corridor_weekly", "Corridor x week x segment", "corridor, week, segment", "Weekly imports (t/week, total cargo) with pre-war base mean; corridor level only"),
    ("model_checks", "Check", "-", "Rebuild of segment totals from rounded inputs vs registered values"),
    ("",),
    ("Security limits (project protocol, section 10): no reserve stocks, no days-to-depletion, no facility-level capacity, routes or single points of failure. All outputs are at corridor level.",),
]


def write_readme_sheet(path):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("README", 0)
    for r in README_ROWS:
        ws.append(list(r))
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, bold=(c.row in (1, 5)))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.row == 5:
                c.fill = PatternFill("solid", fgColor="DDDDDD")
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    for col, wdt in zip("ABCD", [26, 34, 52, 62]):
        ws.column_dimensions[col].width = wdt
    for s in wb.worksheets[1:]:
        for c in s[1]:
            c.font = Font(name="Arial", bold=True)
        s.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inputs", default="data/model_inputs_2026-10-06"); ap.add_argument("--outputs", default="outputs")
    ap.add_argument("--results", default="results"); ap.add_argument("--xlsx", default=None)
    a = ap.parse_args(); os.makedirs(a.results, exist_ok=True)
    d = load(a.inputs); b = d["reg"]["benchmark_food_share"]
    req = subgroup_requirements(d); scen = scenarios(d); cov = coverage(d, scen, b)
    gsub = gaps_by_subgroup(req, cov, d["reg"]); off = off_port(req, d["reg"]); sens = food_share_sensitivity(cov)
    nut = nutrition(d); sup_sum, sup_long = suppliers(d); weekly = corridor_weekly(a.outputs)
    chk = reconstruction_check(d, req)
    commodities = d["sub"].merge(d["cat"][["category_id", "category", "priority", "confidence"]], on="category_id")
    commodities["segment_name"] = commodities.segment.map(SEGMENT_NAMES).fillna("Exposure proxy only")
    demand_levels = pd.DataFrame({"demand": DEMANDS, "description": ["2026 residents, low population factor (1.0618)", "2026 residents, high population factor (1.1863)"], "sort_order": [1, 2]})
    segments = pd.DataFrame([{"segment": k, "segment_name": v, "in_port_model": k in PORT_SEGMENTS, "segment_order": i}
                             for i, (k, v) in enumerate(SEGMENT_NAMES.items(), start=1)])
    tables = {"commodities": commodities, "demand_levels": demand_levels, "segments": segments, "categories": d["cat"], "scenarios": scen,
              "requirement_by_subgroup": req, "corridor_increments": d["inc"].merge(scen, on=["scope", "activity_level"]),
              "coverage_by_segment": cov, "gap_by_subgroup": gsub, "off_port_requirements": off,
              "food_share_sensitivity": sens, "nutrition": nut, "supplier_concentration": sup_sum, "suppliers_top5": sup_long,
              "corridor_weekly": weekly, "model_checks": chk}
    rnd = lambda t: t.apply(lambda c: c.round(6) if pd.api.types.is_float_dtype(c) else c)
    for name, t in tables.items():
        rnd(t).to_csv(os.path.join(a.results, f"{name}.csv"), index=False)
    if a.xlsx:
        os.makedirs(os.path.dirname(a.xlsx), exist_ok=True)
        with pd.ExcelWriter(a.xlsx, engine="openpyxl") as xw:
            for name, t in tables.items():
                rnd(t).to_excel(xw, sheet_name=name[:31], index=False)
        write_readme_sheet(a.xlsx)
    print(f"{len(tables)} tables written to {a.results}")
    return tables


if __name__ == "__main__":
    main()
